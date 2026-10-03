"""Source adapter: profile a real extract, map it, emit the canonical CSVs.

The engine's importer is deliberately strict: two files, fixed columns, a
dataset-wide unique invoice_id. No ERP extract looks like that. This module is
the layer in between. It never relaxes the engine's contract; it produces input
that satisfies it, and records exactly what it did.

Three stages, each inspectable on its own:

    profile_source(data)        -> SourceProfile   what is actually in the file
    infer_mapping(profile, ...) -> SourceMapping   a proposal, for a human to confirm
    apply_mapping(data, mapping)-> PreparedTable   canonical CSV + extras + notes

Design rules, in the spirit of the rest of the engine:

- A mapping is a versioned, fingerprinted artifact. Fold mapping_id into the
  engine snapshot so a changed mapping invalidates approvals, exactly as a
  changed MatchConfig does.
- Inference proposes; it never silently decides anything irreversible. Where
  evidence is genuinely ambiguous (DD/MM vs MM/DD) it refuses and asks.
- Unmapped columns are retained as extras, never dropped and never fed to
  matching. Rejecting a 47-column supplier master is what kills a pilot.
- Every transformation that changes a value is counted and reported.

Standard library only. openpyxl is used for .xlsx if installed, and its absence
is an explicit error rather than a silent skip.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_EVEN
from pathlib import Path

VERSION = "source-adapter-v2"
MAX_SOURCE_BYTES = 20_000_000
MAX_SOURCE_ROWS = 20_000


def sensitive_column(name):
    key = normalize_header(name)
    return key in {"bankn", "iban", "bankaccount", "bankaccountnumber", "accountnumber",
                   "routingnumber", "sortcode", "password", "apikey", "token", "secret"}

# Mirrors engine.SUPPLIER_FIELDS / SPEND_FIELDS plus the optional columns the
# engine accepts. Imported lazily in canonical_fields() to avoid a cycle.
SUPPLIER_OPTIONAL = ("address", "aliases", "lei", "parent_lei", "bank_account_hash")

# ISO 4217 minor units that are not 2. The engine's schema accepts at most two
# decimal places, so a 3-decimal currency has to be reported, not silently cut.
MINOR_UNITS = {
    **{c: 0 for c in ("BIF", "CLP", "DJF", "GNF", "ISK", "JPY", "KMF", "KRW",
                      "PYG", "RWF", "UGX", "UYI", "VND", "VUV", "XAF", "XOF", "XPF")},
    **{c: 3 for c in ("BHD", "IQD", "JOD", "KWD", "LYD", "OMR", "TND")},
    "CLF": 4, "UYW": 4,
}

ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
DELIMITERS = (",", ";", "\t", "|")

# Values operators type into a mandatory field. Kept here as well as in
# governance because profiling must flag them before the engine ever sees them.
PLACEHOLDER_TEXT = {
    "", "-", "--", "---", "----", ".", "..", "N/A", "NA", "N.A.", "NIL", "NONE",
    "NULL", "UNKNOWN", "NOT AVAILABLE", "NOTAVAILABLE", "NOT APPLICABLE",
    "NOTAPPLICABLE", "MISSING", "TBD", "TBC", "TBA", "PENDING", "SAME",
    "AS ABOVE", "ASABOVE", "SEE ABOVE", "DUMMY", "TEMP", "TEST", "VARIOUS",
    "MISC", "SUNDRY", "XXX", "XXXX", "0", "00", "000", "0000", "999", "9999",
}

# Header synonyms. Lowercased, non-alphanumerics stripped, so "Vendor No.",
# "VENDOR_NO" and "vendorno" all collide on one key.
SUPPLIER_SYNONYMS = {
    "supplier_id": ("supplierid", "vendorid", "vendorno", "vendornumber", "lifnr",
                    "suppliernumber", "supplierno", "vendorcode", "suppliercode",
                    "accountnumber", "partnerid", "bpnumber", "creditorid"),
    "name": ("name", "suppliername", "vendorname", "name1", "legalname",
             "companyname", "businessname", "tradingname", "partnername"),
    "country": ("country", "countrycode", "land1", "countrykey", "countryregion",
                "isocountry", "ctry"),
    "registration_id": ("registrationid", "registrationnumber", "companynumber",
                        "companieshouse", "crn", "regno", "registrationno",
                        "legalregistrationnumber", "orgnumber",
                        "organisationnumber", "organizationnumber"),
    "tax_id": ("taxid", "vatnumber", "vatregistrationnumber", "vatno", "vat",
               "stceg", "taxnumber", "taxregistrationnumber", "gstin", "tin",
               "ustid", "siret", "abn"),
    "postcode": ("postcode", "postalcode", "zip", "zipcode", "pstlz", "pincode"),
    "address": ("address", "address1", "addressline1", "street", "stras",
                "streetaddress", "fulladdress", "registeredaddress"),
    "aliases": ("aliases", "alternatenames", "alsoknownas", "aka", "name2",
                "tradingnames", "previousname"),
    "lei": ("lei", "leicode", "legalentityidentifier"),
    "parent_lei": ("parentlei", "parentlegalentityidentifier", "ultimateparentlei"),
    "bank_account_hash": ("bankaccounthash", "bankhash", "accounthash"),
}

SPEND_SYNONYMS = {
    "invoice_id": ("invoiceid", "invoiceno", "invoicenumber", "documentno",
                   "documentnumber", "belnr", "billnumber", "voucherno",
                   "transactionid", "xblnr", "referencenumber"),
    "supplier_id": SUPPLIER_SYNONYMS["supplier_id"],
    "invoice_date": ("invoicedate", "documentdate", "bldat", "postingdate",
                     "budat", "date", "transactiondate", "billdate"),
    "amount": ("amount", "netamount", "net", "amountnet", "invoiceamount",
               "value", "dmbtr", "wrbtr", "totalamount", "linetotal",
               "amountindocumentcurrency"),
    "currency": ("currency", "currencycode", "waers", "curr", "txncurrency",
                 "documentcurrency"),
    "category": ("category", "spendcategory", "commodity", "commoditycode",
                 "matkl", "materialgroup", "glaccount", "costcategory",
                 "expensetype"),
    "description": ("description", "text", "sgtxt", "lineitemtext", "narrative",
                    "details", "itemdescription", "memo"),
}

# Recognised but not part of the engine schema. Captured so the tax model is
# visible in the report even while the engine consumes a single net amount.
AUXILIARY_SYNONYMS = {
    "tax_amount": ("taxamount", "vatamount", "tax", "vat", "wmwst",
                   "mwskzamt", "taxvalue", "salestax", "vatvalue"),
    "gross_amount": ("grossamount", "gross", "amountgross", "totalinclvat",
                     "totalincludingvat", "grosstotal"),
    "document_type": ("documenttype", "blart", "doctype", "invoicetype",
                      "transactiontype"),
    "company_code": ("companycode", "bukrs", "entity", "legalentity", "orgunit"),
}


class IngestError(ValueError):
    """The source cannot be prepared safely without a human decision."""


def _fingerprint(value) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode()
    ).hexdigest()


def normalize_header(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value).casefold())
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in text if not unicodedata.combining(c)))


def canonical_fields(target: str) -> tuple[str, ...]:
    from .engine import SPEND_FIELDS, SUPPLIER_FIELDS
    if target == "suppliers":
        return SUPPLIER_FIELDS
    if target == "spend":
        return SPEND_FIELDS
    raise IngestError("target must be 'suppliers' or 'spend'")


# --------------------------------------------------------------------------
# Stage 1: profiling
# --------------------------------------------------------------------------

@dataclass
class ColumnProfile:
    name: str
    index: int
    filled: int
    total: int
    distinct: int
    samples: list
    inferred_type: str
    placeholder_count: int
    placeholder_values: list
    max_length: int

    @property
    def fill_rate(self) -> float:
        return self.filled / self.total if self.total else 0.0

    @property
    def unique(self) -> bool:
        return self.filled == self.total and self.distinct == self.total and self.total > 0


@dataclass
class SourceProfile:
    name: str
    encoding: str
    delimiter: str
    header_row: int
    row_count: int
    columns: list
    notes: list = field(default_factory=list)
    content_sha256: str = ""

    def column(self, name: str):
        for item in self.columns:
            if item.name == name:
                return item
        return None

    def report(self, include_samples=True) -> dict:
        return {
            "version": VERSION, "name": self.name, "encoding": self.encoding,
            "delimiter": self.delimiter, "header_row": self.header_row,
            "row_count": self.row_count, "content_sha256": self.content_sha256,
            "notes": self.notes,
            "columns": [{
                "name": c.name, "index": c.index, "fill_rate": round(c.fill_rate, 4),
                "distinct": c.distinct, "unique": c.unique, "type": c.inferred_type,
                "max_length": c.max_length, "placeholders": c.placeholder_count,
                "placeholder_values": [] if sensitive_column(c.name) else c.placeholder_values,
                "samples": (["[REDACTED]"] if sensitive_column(c.name) and c.filled else c.samples) if include_samples else [],
            } for c in self.columns],
        }


def decode(data) -> tuple[str, str]:
    """Return (text, encoding). Tries UTF-8 first, then the usual ERP fallbacks."""
    if isinstance(data, str):
        return data, "utf-8"
    if not isinstance(data, (bytes, bytearray)):
        raise IngestError("Provide text or bytes")
    for encoding in ENCODINGS:
        try:
            return bytes(data).decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise IngestError("Could not decode the file as UTF-8, CP1252 or Latin-1")


def sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:50])
    best, best_score = ",", -1.0
    for candidate in DELIMITERS:
        counts = [line.count(candidate) for line in sample.splitlines() if line.strip()]
        if not counts or max(counts) == 0:
            continue
        common = Counter(counts).most_common(1)[0]
        # Prefer the delimiter with the most columns, consistently present.
        score = common[0] * (common[1] / len(counts))
        if score > best_score:
            best, best_score = candidate, score
    return best


def find_header_row(rows: list) -> int:
    """First full-width row whose cells are non-empty, distinct and non-numeric.

    ERP exports routinely carry title and filter lines above the real header.
    Those banner lines are narrower than the data, so the modal row width tells
    us how wide a real row is and the banners fall away on their own.
    """
    widths = Counter(len([c for c in row if str(c).strip()]) or len(row) for row in rows)
    if not widths:
        return 0
    modal = max(widths, key=lambda w: (widths[w], w))
    for index, row in enumerate(rows[:20]):
        cells = [str(c).strip() for c in row]
        if len(cells) < max(2, modal) or any(not c for c in cells[:modal]):
            continue
        if len({normalize_header(c) for c in cells[:modal]}) != modal:
            continue
        numeric = sum(bool(re.fullmatch(r"-?[\d.,]+", c)) for c in cells[:modal])
        if numeric <= modal // 3:
            return index
    return 0


def _infer_type(values: list) -> str:
    sample = [v for v in values if v][:400]
    if not sample:
        return "empty"
    def ratio(predicate):
        return sum(bool(predicate(v)) for v in sample) / len(sample)
    if ratio(lambda v: parse_date(v, strict=False) is not None) > .9:
        return "date"
    if ratio(lambda v: re.fullmatch(r"[A-Za-z]{3}", v)) > .9:
        return "currency_code"
    if ratio(lambda v: re.fullmatch(r"[A-Za-z]{2}", v)) > .9:
        return "country_code"
    if ratio(lambda v: _looks_numeric(v)) > .9:
        return "number"
    if ratio(lambda v: re.fullmatch(r"\d+", v)) > .9:
        return "integer"
    return "text"


def _looks_numeric(value: str) -> bool:
    cleaned = re.sub(r"[\s '’]", "", str(value))
    cleaned = re.sub(r"^[^\d\-+(]*", "", cleaned)
    cleaned = re.sub(r"[A-Za-z\s]*$", "", cleaned)
    return bool(re.fullmatch(r"[-+(]?\d{1,3}(?:[.,]\d{3})*(?:[.,]\d+)?\)?|[-+(]?\d+(?:[.,]\d+)?\)?", cleaned))


def read_table(data, name="source", delimiter=None, header_row=None) -> tuple[list, list, SourceProfile]:
    """Return (header, rows, partial profile). Rows are lists of strings."""
    if isinstance(data, Path) or (
            isinstance(data, str) and "\n" not in data and Path(data).suffix.lower() in {".xlsx", ".xlsm"}):
        if Path(data).suffix.lower() not in {".xlsx", ".xlsm"}:
            return read_table(_load(data), name, delimiter, header_row)
        return _read_xlsx(Path(data), name, header_row)
    text, encoding = decode(data)
    if not text.strip():
        raise IngestError(f"{name}: the file is empty")
    if len(text.encode("utf-8")) > MAX_SOURCE_BYTES or "\x00" in text:
        raise IngestError(f"{name}: source exceeds 20 MB or contains null bytes")
    sha = hashlib.sha256(bytes(data) if isinstance(data, (bytes, bytearray)) else text.encode("utf-8")).hexdigest()
    delimiter = delimiter or sniff_delimiter(text)
    try:
        raw = []
        for row in csv.reader(io.StringIO(text.lstrip("﻿"), newline=""), delimiter=delimiter, strict=True):
            raw.append(row)
            if len(raw) > MAX_SOURCE_ROWS + 20:
                raise IngestError(f"{name}: source exceeds {MAX_SOURCE_ROWS} rows")
    except csv.Error as exc:
        raise IngestError(f"{name}: malformed CSV") from exc
    if not raw or not any(any(c.strip() for c in row) for row in raw):
        raise IngestError(f"{name}: no data rows found")
    index = find_header_row(raw) if header_row is None else header_row
    if isinstance(index, bool) or not isinstance(index, int) or not 0 <= index < len(raw):
        raise IngestError(f"{name}: invalid header row")
    header = _validate_header(raw[index], name)
    width = len(header)
    rows = _validate_rows(raw[index + 1:], width, name)
    profile = SourceProfile(name=name, encoding=encoding, delimiter=delimiter,
                            header_row=index, row_count=len(rows), columns=[],
                            content_sha256=sha)
    if index > 0:
        profile.notes.append(f"Skipped {index} line(s) above the header row.")
    return header, rows, profile


def _validate_header(row, name):
    header = [str(c).strip() for c in row]
    keys = [normalize_header(c) for c in header]
    if not all(keys) or len(keys) != len(set(keys)):
        raise IngestError(f"{name}: empty or duplicate header; specify a correct header row")
    return header


def _validate_rows(rows, width, name):
    result = []
    for row in rows:
        if not any(c.strip() for c in row):
            continue
        if len(row) != width:
            raise IngestError(f"{name}: row width differs from header; no columns were discarded")
        if any(len(c) > 1000 for c in row):
            raise IngestError(f"{name}: field exceeds 1000 characters")
        result.append(row)
    if not result or len(result) > MAX_SOURCE_ROWS:
        raise IngestError(f"{name}: require 1..{MAX_SOURCE_ROWS} source rows")
    return result


def _read_xlsx(path: Path, name: str, header_row=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise IngestError("Reading .xlsx needs openpyxl: python -m pip install openpyxl") from exc
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise IngestError(f"{name}: workbook exceeds 20 MB")
    contents = path.read_bytes()
    book = load_workbook(io.BytesIO(contents), read_only=True, data_only=False)
    sheet = book[book.sheetnames[0]]
    raw = []
    try:
        for row in sheet.iter_rows():
            if any(c.data_type == "f" for c in row):
                raise IngestError(f"{name}: export formula results to values before import")
            cells = []
            for cell in row:
                value = cell.value
                if isinstance(value, (float, int)) and re.fullmatch(r"0{2,}", cell.number_format):
                    raise IngestError(f"{name}: numeric zero-padded identifiers must be exported as text")
                cells.append("" if value is None else (value.strftime("%Y-%m-%d") if isinstance(value, (datetime, date)) else str(value)))
            raw.append(cells)
            if len(raw) > MAX_SOURCE_ROWS + 20:
                raise IngestError(f"{name}: workbook exceeds {MAX_SOURCE_ROWS} rows")
    finally:
        book.close()
    if not raw:
        raise IngestError(f"{name}: the first worksheet is empty")
    index = find_header_row(raw) if header_row is None else header_row
    if isinstance(index, bool) or not isinstance(index, int) or not 0 <= index < len(raw):
        raise IngestError(f"{name}: invalid header row")
    header = _validate_header(raw[index], name)
    width = len(header)
    rows = _validate_rows(raw[index + 1:], width, name)
    profile = SourceProfile(name=name, encoding="xlsx", delimiter="", header_row=index,
                            row_count=len(rows), columns=[],
                            content_sha256=hashlib.sha256(contents).hexdigest())
    profile.notes.append(f"Read worksheet {sheet.title!r}; other sheets ignored.")
    return header, rows, profile


def profile_source(data, name="source", delimiter=None, header_row=None) -> SourceProfile:
    header, rows, profile = read_table(data, name, delimiter, header_row)
    seen = Counter()
    for position, raw_name in enumerate(header):
        label = raw_name or f"column_{position + 1}"
        seen[label] += 1
        if seen[label] > 1:
            label = f"{label}__{seen[label]}"
            profile.notes.append(f"Duplicate header {raw_name!r} renamed to {label!r}.")
        values = [r[position].strip() for r in rows]
        filled = [v for v in values if v]
        placeholders = [v for v in filled if v.upper() in PLACEHOLDER_TEXT]
        profile.columns.append(ColumnProfile(
            name=label, index=position, filled=len(filled), total=len(values),
            distinct=len(set(filled)), samples=list(dict.fromkeys(filled))[:5],
            inferred_type=_infer_type(values), placeholder_count=len(placeholders),
            placeholder_values=sorted({v for v in placeholders})[:5],
            max_length=max((len(v) for v in filled), default=0)))
    for column in profile.columns:
        if column.placeholder_count:
            profile.notes.append(
                f"{column.name}: {column.placeholder_count} placeholder value(s) "
                f"({', '.join(column.placeholder_values)}) — these are not identity.")
    return profile


# --------------------------------------------------------------------------
# Value transforms
# --------------------------------------------------------------------------

DATE_PATTERNS = (
    (re.compile(r"^(\d{4})-(\d{2})-(\d{2})$"), "ymd"),
    (re.compile(r"^(\d{8})$"), "compact"),
    (re.compile(r"^(\d{4})/(\d{2})/(\d{2})$"), "ymd"),
    (re.compile(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})$"), "ambiguous"),
    (re.compile(r"^(\d{1,2})[\- ]([A-Za-z]{3})[\- ](\d{4})$"), "dmonthy"),
)
MONTHS = {m: i for i, m in enumerate(
    ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}


def parse_date(value: str, order="dmy", strict=True):
    """Return an ISO date string, or None. order applies only to ambiguous forms."""
    text = str(value).strip()
    if not text:
        return None
    text = text.split("T")[0].split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2}[T ]", text) else text
    for pattern, kind in DATE_PATTERNS:
        match = pattern.match(text)
        if not match:
            continue
        try:
            if kind == "compact":
                digits = match.group(1)
                # SAP writes YYYYMMDD. Fall back to the column's order when
                # those first four digits cannot be a year and month and day.
                try:
                    return date(int(digits[:4]), int(digits[4:6]), int(digits[6:])).isoformat()
                except ValueError:
                    pass
                a, b, y = int(digits[:2]), int(digits[2:4]), int(digits[4:])
                d, m = (a, b) if order == "dmy" else (b, a)
            elif kind == "ymd":
                y, m, d = (int(g) for g in match.groups())
            elif kind == "dmonthy":
                d, month, y = match.group(1), match.group(2).lower()[:3], match.group(3)
                if month not in MONTHS:
                    return None
                d, m, y = int(d), MONTHS[month], int(y)
            else:
                a, b, y = (int(g) for g in match.groups())
                d, m = (a, b) if order == "dmy" else (b, a)
            return date(y, m, d).isoformat()
        except ValueError:
            return None
    return None


def detect_date_order(values) -> str:
    """'dmy', 'mdy', 'ymd', 'ambiguous' or 'conflicting'.

    'ymd' means no value in the column can be read two ways, so nothing needs
    deciding. 'ambiguous' means some value could be either and the data does
    not settle it — the caller must ask rather than guess.
    """
    dmy = mdy = False
    ambiguous_values = 0
    for value in values:
        text = str(value).strip()
        match = re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})[/.\-]\d{4}", text)
        if match:
            first, second = int(match.group(1)), int(match.group(2))
        elif re.fullmatch(r"\d{8}", text):
            try:  # An unambiguous YYYYMMDD carries no day/month evidence.
                date(int(text[:4]), int(text[4:6]), int(text[6:]))
                continue
            except ValueError:
                first, second = int(text[:2]), int(text[2:4])
        else:
            continue
        ambiguous_values += 1
        if first > 12:
            dmy = True
        if second > 12:
            mdy = True
    if not ambiguous_values:
        return "ymd"
    if dmy and not mdy:
        return "dmy"
    if mdy and not dmy:
        return "mdy"
    if not dmy and not mdy:
        return "ambiguous"
    return "conflicting"


def detect_decimal_style(values) -> str:
    """'point', 'comma' or 'plain'. Decided on evidence, never on locale guesses."""
    comma = point = 0
    uncertain = False
    for value in values:
        text = re.sub(r"[^\d.,]", "", str(value))
        if re.fullmatch(r"\d{1,3}(?:\.\d{3})+,\d{1,}", text):
            comma += 1
        elif re.fullmatch(r"\d{1,3}(?:,\d{3})+\.\d{1,}", text):
            point += 1
        elif re.fullmatch(r"\d+,\d{1,2}", text):
            comma += 1
        elif re.fullmatch(r"\d+\.\d{1,2}", text):
            point += 1
        elif re.search(r"[.,]", text):
            uncertain = True
    if comma and point:
        return "conflicting"
    if comma and not point:
        return "comma"
    if point and not comma:
        return "point"
    return "ambiguous" if uncertain else "plain"


def parse_amount(value: str, style="point"):
    """Return a Decimal, or None. Handles (1,234.56), 1.234,56, 1 234,56, 100 CR."""
    text = str(value).strip()
    if not text:
        return None
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
    trailing = re.search(r"\b(CR|DR)\b\s*$", text, re.IGNORECASE)
    if trailing:
        negative = negative or trailing.group(1).upper() == "CR"
        text = text[:trailing.start()]
    # Accept a currency symbol or code, never arbitrary text around digits.
    text = re.sub(r"^[£€$¥]\s*|\s*[£€$¥]$", "", text.strip())
    text = re.sub(r"^[A-Z]{3}\s+|\s+[A-Z]{3}$", "", text)
    text = re.sub(r"[\s '’]", "", text)
    if text.startswith(("-", "+")):
        if negative:
            return None
        negative, text = text.startswith("-"), text[1:]
    if not text:
        return None
    if style == "comma":
        if not re.fullmatch(r"(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?", text):
            return None
        text = text.replace(".", "").replace(",", ".")
    elif style == "point":
        if not re.fullmatch(r"(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", text):
            return None
        text = text.replace(",", "")
    else:
        if not re.fullmatch(r"\d+", text):
            return None
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    return -amount if negative else amount


def strip_leading_zeros(value: str) -> str:
    text = str(value).strip()
    stripped = text.lstrip("0")
    return stripped if stripped else ("0" if text else "")


TRANSFORMS = {
    "trim": lambda v, _: str(v).strip(),
    "upper": lambda v, _: str(v).strip().upper(),
    "strip_leading_zeros": lambda v, _: strip_leading_zeros(v),
    "collapse_space": lambda v, _: re.sub(r"\s+", " ", str(v)).strip(),
    "blank_placeholders": lambda v, _: "" if str(v).strip().upper() in PLACEHOLDER_TEXT else str(v).strip(),
}


# --------------------------------------------------------------------------
# Stage 2: mapping
# --------------------------------------------------------------------------

@dataclass(frozen=True)
class SourceMapping:
    source_system: str
    target: str                      # "suppliers" | "spend"
    columns: dict                    # canonical field -> source column name
    transforms: dict = field(default_factory=dict)   # field -> [transform names]
    auxiliary: dict = field(default_factory=dict)    # tax_amount/gross_amount/... -> column
    date_order: str = "dmy"
    decimal_style: str = "point"
    surrogate_keys: bool = True      # prefix native IDs with the source system
    invoice_key_scope: str = "dataset"   # "supplier" when numbers repeat per vendor
    key_separator: str = ":"
    passthrough: bool = True         # keep unmapped columns as extras
    allow_rejected_rows: bool = False   # an explicit mapping policy, never inference
    version: str = VERSION

    def __post_init__(self):
        if self.version != VERSION:
            raise IngestError("Unsupported mapping version; regenerate and review the artifact")
        for value in (self.surrogate_keys, self.passthrough, self.allow_rejected_rows):
            if not isinstance(value, bool):
                raise IngestError("Mapping policy switches must be boolean")
        if self.target not in {"suppliers", "spend"}:
            raise IngestError("target must be 'suppliers' or 'spend'")
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,40}", self.source_system):
            raise IngestError("source_system must be 1..40 letters, digits, underscores or hyphens")
        if self.date_order not in {"dmy", "mdy", "ymd"}:
            raise IngestError("date_order must be dmy, mdy or ymd")
        if self.decimal_style not in {"point", "comma", "plain"}:
            raise IngestError("decimal_style must be point, comma or plain")
        if self.invoice_key_scope not in {"dataset", "supplier"}:
            raise IngestError("invoice_key_scope must be 'dataset' or 'supplier'")
        if self.key_separator not in {":", "|", "~", "/"}:
            raise IngestError("key_separator must be one of : | ~ /")
        required = set(canonical_fields(self.target))
        missing = required - set(self.columns) - {"category", "description", "tax_id", "registration_id", "postcode"}
        if missing:
            raise IngestError("Mapping is missing required field(s): " + ", ".join(sorted(missing)))
        allowed = required | (set(SUPPLIER_OPTIONAL) if self.target == "suppliers" else set())
        if set(self.columns) - allowed or set(self.transforms) - allowed or set(self.auxiliary) - set(AUXILIARY_SYNONYMS):
            raise IngestError("Unknown canonical or auxiliary field in mapping")
        if any(not isinstance(c, str) or not c for c in (*self.columns.values(), *self.auxiliary.values())):
            raise IngestError("Mapped columns must be non-empty names")
        for names in self.transforms.values():
            unknown = set(names) - set(TRANSFORMS)
            if unknown:
                raise IngestError("Unknown transform(s): " + ", ".join(sorted(unknown)))

    def canonical(self) -> str:
        return json.dumps(asdict(self), sort_keys=True, separators=(",", ":"))

    @property
    def mapping_id(self) -> str:
        return hashlib.sha256(self.canonical().encode()).hexdigest()

    def to_json(self, indent=2) -> str:
        return json.dumps({**asdict(self), "mapping_id": self.mapping_id}, indent=indent, sort_keys=True)

    @classmethod
    def from_json(cls, text: str) -> "SourceMapping":
        data = json.loads(text)
        if not isinstance(data, dict):
            raise IngestError("Mapping must be a JSON object")
        declared = data.pop("mapping_id", None)
        try:
            mapping = cls(**data)
        except (TypeError, AttributeError) as exc:
            raise IngestError("Invalid mapping structure") from exc
        if declared is not None and declared != mapping.mapping_id:
            raise IngestError("Mapping fingerprint does not match its contents")
        return mapping


def infer_mapping(profile: SourceProfile, target: str, source_system="SRC", data=None, **overrides) -> SourceMapping:
    """Propose a mapping from header synonyms and column evidence.

    This is a proposal for a human to confirm. Pass `data` (the same bytes the
    profile was built from) so date order and decimal style can be decided on
    evidence; without it, those must be supplied explicitly. It refuses rather
    than guesses when a date column could be read two ways.
    """
    synonyms = SUPPLIER_SYNONYMS if target == "suppliers" else SPEND_SYNONYMS
    by_key = {}
    for column in profile.columns:
        by_key.setdefault(normalize_header(column.name), column)
    columns, unresolved = {}, []
    for canonical_name, candidates in synonyms.items():
        for candidate in candidates:
            if candidate in by_key:
                columns[canonical_name] = by_key[candidate].name
                break
        else:
            if canonical_name in canonical_fields(target):
                unresolved.append(canonical_name)
    auxiliary = {}
    for name, candidates in AUXILIARY_SYNONYMS.items():
        for candidate in candidates:
            if candidate in by_key:
                auxiliary[name] = by_key[candidate].name
                break
    columns.update(overrides.pop("column_overrides", {}))
    if target == "suppliers" and any(normalize_header(c.name) == "stcd1" for c in profile.columns):
        profile.notes.append("STCD1 is a tax-number field, not automatic legal registration evidence. Map it explicitly only under a verified source contract.")

    # Fall back to column evidence for the fields that matter most.
    if target == "suppliers" and "country" not in columns:
        for column in profile.columns:
            if column.inferred_type == "country_code":
                columns["country"] = column.name
                break
    if target == "spend":
        if "currency" not in columns:
            for column in profile.columns:
                if column.inferred_type == "currency_code":
                    columns["currency"] = column.name
                    break
        if "invoice_date" not in columns:
            for column in profile.columns:
                if column.inferred_type == "date":
                    columns["invoice_date"] = column.name
                    break
    if target == "suppliers" and "supplier_id" not in columns:
        for column in profile.columns:
            if column.unique and column.inferred_type in {"integer", "text"}:
                columns["supplier_id"] = column.name
                break

    transforms = {}
    # Native IDs keep leading zeros. Stripping them requires an explicit mapping
    # transform shared by both sources, otherwise references can collide.
    for field_name in ("supplier_id", "invoice_id"):
        if field_name in columns:
            transforms[field_name] = ["trim"]
    for field_name in ("registration_id", "tax_id", "lei", "parent_lei"):
        if field_name in columns:
            transforms[field_name] = ["blank_placeholders", "upper"]
    if "name" in columns:
        transforms["name"] = ["collapse_space"]
    for field_name in ("country", "currency"):
        if field_name in columns:
            transforms[field_name] = ["trim", "upper"]

    date_order, decimal_style = "ymd", "point"
    if target == "spend" and data is not None:
        header, rows, _ = read_table(data, profile.name, profile.delimiter or None, profile.header_row)
        position = {name: i for i, name in enumerate(header)}
        if "invoice_date" in columns:
            values = [r[position[columns["invoice_date"]]] for r in rows] if columns["invoice_date"] in position else []
            detected = detect_date_order(values)
            if detected == "conflicting":
                raise IngestError(
                    f"{profile.name}: the date column {columns['invoice_date']!r} contains both "
                    "DD/MM and MM/DD values. Split the source or set date_order explicitly.")
            if detected == "ambiguous" and "date_order" not in overrides:  # refuse, do not guess
                raise IngestError(
                    f"{profile.name}: the date column {columns['invoice_date']!r} is ambiguous "
                    "(no value has a day above 12). Pass date_order='dmy' or 'mdy' explicitly.")
            date_order = detected if detected in {"dmy", "mdy", "ymd"} else "dmy"
        if "amount" in columns and columns["amount"] in position:
            decimal_style = detect_decimal_style([r[position[columns["amount"]]] for r in rows])
            if decimal_style in {"ambiguous", "conflicting"} and "decimal_style" not in overrides:
                raise IngestError(f"{profile.name}: {decimal_style} decimal conventions; set decimal_style explicitly or split the source")
            if decimal_style in {"ambiguous", "conflicting"}:
                decimal_style = overrides["decimal_style"]
        if {"invoice_id", "supplier_id"} <= set(columns):
            numbers = [r[position[columns["invoice_id"]]].strip() for r in rows]
            pairs = list(zip((r[position[columns["supplier_id"]]].strip() for r in rows), numbers))
            if len(set(pairs)) > len(set(numbers)):
                overrides.setdefault("invoice_key_scope", "supplier")

    mapping = SourceMapping(source_system=source_system, target=target, columns=columns,
                            transforms=transforms, auxiliary=auxiliary,
                            date_order=overrides.pop("date_order", date_order),
                            decimal_style=overrides.pop("decimal_style", decimal_style),
                            **overrides)
    if unresolved:
        profile.notes.append("Unmapped required field(s), confirm by hand: " + ", ".join(sorted(unresolved)))
    return mapping




# --------------------------------------------------------------------------
# Stage 3: application
# --------------------------------------------------------------------------

@dataclass
class PreparedTable:
    csv_text: str
    extras: dict                 # canonical record id -> {source column: value}
    notes: list
    rejected: list               # rows dropped, with a reason each
    mapping_id: str
    row_count: int


def _apply_transforms(value, names):
    for name in names or ():
        value = TRANSFORMS[name](value, None)
    return value


def apply_mapping(data, mapping: SourceMapping, name=None, profile: SourceProfile | None = None) -> PreparedTable:
    """Produce a canonical CSV the engine's importer accepts without changes."""
    label = name or (profile.name if profile else mapping.target)
    header, rows, actual_profile = read_table(data, label,
                                 delimiter=profile.delimiter if profile and profile.delimiter else None,
                                 header_row=profile.header_row if profile else None)
    if profile and actual_profile.content_sha256 != profile.content_sha256:
        raise IngestError(f"{label}: source changed after profiling; profile and confirm again")
    position = {column: index for index, column in enumerate(header)}
    for field_name, column in {**mapping.columns, **mapping.auxiliary}.items():
        if column not in position:
            raise IngestError(f"{label}: mapped column {column!r} for {field_name} is not in this file")

    fields = canonical_fields(mapping.target)
    optional = SUPPLIER_OPTIONAL if mapping.target == "suppliers" else ()
    emitted = [f for f in (*fields, *optional) if f in mapping.columns or f in fields]
    mapped_columns = set(mapping.columns.values()) | set(mapping.auxiliary.values())
    extra_columns = [c for c in header if c not in mapped_columns] if mapping.passthrough else []

    notes, rejected, extras, records = [], [], {}, []
    counters = Counter()
    minor_unit_hits = Counter()

    def surrogate(raw_value):
        text = str(raw_value).strip()
        if not mapping.surrogate_keys or not text:
            return text
        return f"{mapping.source_system}{mapping.key_separator}{text}"

    seen_records = {}
    for line, row in enumerate(rows, start=actual_profile.header_row + 2):
        record = {}
        for field_name in emitted:
            column = mapping.columns.get(field_name)
            raw = row[position[column]] if column else ""
            value = _apply_transforms(raw, mapping.transforms.get(field_name))
            if value != str(raw).strip():
                counters[field_name] += 1
            record[field_name] = value

        native = {k: record.get(k, "") for k in ("supplier_id", "invoice_id") if k in record}
        clashing = next((k for k, v in native.items() if v and mapping.key_separator in v), None)
        if clashing:
            rejected.append({"line": line,
                             "reason": f"{clashing} contains the key separator {mapping.key_separator!r}"})
            continue
        if native.get("supplier_id"):
            record["supplier_id"] = surrogate(native["supplier_id"])
        if native.get("invoice_id"):
            # An invoice number unique only within a vendor becomes unique
            # dataset-wide by carrying its vendor, rather than being rejected.
            if mapping.target == "spend" and mapping.invoice_key_scope == "supplier" and native.get("supplier_id"):
                record["invoice_id"] = surrogate(
                    native["supplier_id"] + mapping.key_separator + native["invoice_id"])
            else:
                record["invoice_id"] = surrogate(native["invoice_id"])

        if mapping.target == "spend":
            iso = parse_date(record.get("invoice_date", ""), mapping.date_order)
            if iso is None:
                rejected.append({"line": line, "reason": f"unparseable date {record.get('invoice_date', '')!r}"})
                continue
            if iso != record["invoice_date"]:
                counters["invoice_date"] += 1
            record["invoice_date"] = iso

            net = parse_amount(record.get("amount", ""), mapping.decimal_style)
            if net is None:
                rejected.append({"line": line, "reason": f"unparseable amount {record.get('amount', '')!r}"})
                continue
            currency = (record.get("currency") or "").upper()
            units = MINOR_UNITS.get(currency, 2)
            if units > 2:
                minor_unit_hits[currency] += 1
            if abs(net) >= Decimal("1000000000000"):
                raise IngestError(f"{label}:{line}: amount exceeds the engine's 12-digit limit")
            try:
                quantized = net.quantize(Decimal("0.01"), rounding=ROUND_HALF_EVEN)
            except InvalidOperation as exc:
                raise IngestError(f"{label}:{line}: amount precision exceeds the current schema") from exc
            if quantized != net or (units == 0 and net != net.to_integral_value()):
                raise IngestError(f"{label}:{line}: {currency} amount cannot be represented exactly by the current engine schema; no rounding performed")
            record["amount"] = f"{quantized:.2f}"

            tax = parse_amount(row[position[mapping.auxiliary["tax_amount"]]], mapping.decimal_style) if "tax_amount" in mapping.auxiliary else None
            gross = parse_amount(row[position[mapping.auxiliary["gross_amount"]]], mapping.decimal_style) if "gross_amount" in mapping.auxiliary else None
            if tax is not None and gross is not None and (net + tax - gross).copy_abs() > Decimal("0.01"):
                counters["tax_reconciliation_failed"] += 1
            for field_name, parsed in (("tax_amount", tax), ("gross_amount", gross)):
                if field_name in mapping.auxiliary and row[position[mapping.auxiliary[field_name]]].strip() and parsed is None:
                    raise IngestError(f"{label}:{line}: invalid auxiliary {field_name}")

        key = record.get("invoice_id") or record.get("supplier_id") or f"row-{line}"
        if key in seen_records and seen_records[key] != row:
            raise IngestError(f"{label}:{line}: conflicting rows share a mapped record ID; choose the correct key scope or transforms")
        seen_records[key] = row
        if extra_columns or mapping.auxiliary:
            # Auxiliary amounts stay in local evidence too; do not discard them
            # merely because the core ledger currently stores one net amount.
            extras[key] = {c: row[position[c]] for c in [*extra_columns, *mapping.auxiliary.values()] if row[position[c]].strip()}
        records.append(record)

    if not records:
        raise IngestError(f"{label}: no rows survived mapping; see the rejected list")
    if rejected and not mapping.allow_rejected_rows:
        raise IngestError(f"{label}: {len(rejected)} rejected row(s); whole source refused. First rejection: {rejected[0]['reason']}. Set allow_rejected_rows explicitly in a reviewed mapping to prepare a partial source")

    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=emitted, lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    writer.writerows(records)

    for field_name, count in sorted(counters.items()):
        notes.append(f"{field_name}: {count} value(s) normalised by the adapter.")
    for currency, count in sorted(minor_unit_hits.items()):
        notes.append(f"{currency} has {MINOR_UNITS[currency]} minor units; {count} row(s) are exactly representable in the current 2-place schema. No rounding performed.")
    if extra_columns:
        notes.append(f"{len(extra_columns)} unmapped column(s) retained as extras, excluded from matching: "
                     + ", ".join(extra_columns[:12]) + ("…" if len(extra_columns) > 12 else ""))
    if rejected:
        notes.append(f"{len(rejected)} row(s) rejected; the rest were prepared.")

    return PreparedTable(csv_text=buffer.getvalue(), extras=extras, notes=notes,
                         rejected=rejected, mapping_id=mapping.mapping_id, row_count=len(records))


def mapping_header_offset(profile) -> int:
    return profile.header_row if profile else 0


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------

@dataclass
class PreparedDataset:
    suppliers_csv: str
    spend_csv: str
    extras: dict
    manifest: dict

    def analyze_with(self, engine, contract_workdir=None):
        """Import into an Engine. The manifest travels as the contract manifest,
        so mapping_id is folded into the snapshot exactly like MatchConfig."""
        validate_prepared(self)
        contract, normalized = self.manifest, None
        if contract_workdir:
            from .contracts import run_dbt_contracts
            normalized, dbt = run_dbt_contracts(self.suppliers_csv, self.spend_csv, contract_workdir)
            contract = {"source_adapter": self.manifest, "dbt": dbt}
        return engine.analyze(self.suppliers_csv, self.spend_csv, normalized, contract)


def prepare_dataset(supplier_source, spend_source, supplier_mapping=None, spend_mapping=None,
                    source_system="SRC", supplier_name="suppliers", spend_name="spend",
                    **inference) -> PreparedDataset:
    """Profile both files, infer or accept mappings, and emit canonical CSVs."""
    supplier_bytes = _load(supplier_source)
    spend_bytes = _load(spend_source)
    supplier_profile = profile_source(supplier_bytes, supplier_name)
    spend_profile = profile_source(spend_bytes, spend_name)
    supplier_mapping = supplier_mapping or infer_mapping(
        supplier_profile, "suppliers", source_system, supplier_bytes, **inference.get("suppliers", {}))
    spend_mapping = spend_mapping or infer_mapping(
        spend_profile, "spend", source_system, spend_bytes, **inference.get("spend", {}))

    if supplier_mapping.target != "suppliers" or spend_mapping.target != "spend":
        raise IngestError("Supplier and spend mappings must target the correct table")
    if supplier_mapping.surrogate_keys != spend_mapping.surrogate_keys or \
            supplier_mapping.source_system != spend_mapping.source_system or \
            supplier_mapping.key_separator != spend_mapping.key_separator:
        raise IngestError("Supplier and spend mappings must share source_system, "
                          "surrogate_keys and key_separator or references will not join")

    suppliers = apply_mapping(supplier_bytes, supplier_mapping, supplier_name, supplier_profile)
    spend = apply_mapping(spend_bytes, spend_mapping, spend_name, spend_profile)

    manifest = {
        "version": VERSION,
        "source_system": supplier_mapping.source_system,
        "mappings": {"suppliers": json.loads(supplier_mapping.to_json()), "spend": json.loads(spend_mapping.to_json())},
        "suppliers": {"mapping_id": suppliers.mapping_id, "rows": suppliers.row_count,
                      "canonical_sha256": hashlib.sha256(suppliers.csv_text.encode()).hexdigest(),
                      "source_sha256": supplier_profile.content_sha256,
                      "rejected": len(suppliers.rejected), "notes": suppliers.notes},
        "spend": {"mapping_id": spend.mapping_id, "rows": spend.row_count,
                  "canonical_sha256": hashlib.sha256(spend.csv_text.encode()).hexdigest(),
                  "source_sha256": spend_profile.content_sha256,
                  "rejected": len(spend.rejected), "notes": spend.notes},
        "profiles": {"suppliers": supplier_profile.report(include_samples=False), "spend": spend_profile.report(include_samples=False)},
        "rejections": {"suppliers": suppliers.rejected, "spend": spend.rejected},
        "limitations": ("Mapped and normalised by the source adapter. Extras are retained "
                        "for evidence but never used for identity matching. Column mapping "
                        "is an operator assertion, not a verified source contract."),
    }
    manifest["manifest_id"] = _fingerprint({k: v for k, v in manifest.items() if k != "profiles"})
    return PreparedDataset(suppliers_csv=suppliers.csv_text, spend_csv=spend.csv_text,
                           extras={"suppliers": suppliers.extras, "spend": spend.extras},
                           manifest=manifest)


def _load(source):
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, Path):
        if source.suffix.lower() in {".xlsx", ".xlsm"}:
            return source
        if source.stat().st_size > MAX_SOURCE_BYTES:
            raise IngestError("Source file exceeds 20 MB")
        return source.read_bytes()
    if isinstance(source, str) and "\n" not in source and Path(source).exists():
        return _load(Path(source))
    if isinstance(source, str):
        return source.encode("utf-8")
    raise IngestError("Provide a path, text or bytes")


def validate_prepared(prepared):
    manifest = prepared.manifest
    if manifest.get("version") != VERSION or manifest.get("manifest_id") != _fingerprint(
            {k: v for k, v in manifest.items() if k not in {"profiles", "manifest_id"}}):
        raise IngestError("Prepared manifest fingerprint or version is invalid")
    for target, text in (("suppliers", prepared.suppliers_csv), ("spend", prepared.spend_csv)):
        mapping = SourceMapping.from_json(json.dumps(manifest["mappings"][target]))
        if mapping.mapping_id != manifest[target]["mapping_id"] or mapping.target != target:
            raise IngestError("Prepared mapping does not match its manifest")
        if hashlib.sha256(text.encode()).hexdigest() != manifest[target]["canonical_sha256"]:
            raise IngestError("Prepared CSV changed after mapping; prepare it again")


def load_prepared(directory):
    root = Path(directory)
    prepared = PreparedDataset((root / "suppliers.csv").read_text(encoding="utf-8"),
                               (root / "spend.csv").read_text(encoding="utf-8"), {},
                               json.loads((root / "manifest.json").read_text(encoding="utf-8")))
    validate_prepared(prepared)
    return prepared


# --------------------------------------------------------------------------
# CLI: python -m enterprise_ai.ingest <profile|map|prepare>
# --------------------------------------------------------------------------

def main(argv=None):
    parser = argparse.ArgumentParser(prog="enterprise_ai.ingest",
                                     description="Profile and map a real extract into the engine's canonical CSVs")
    sub = parser.add_subparsers(dest="command", required=True)

    prof = sub.add_parser("profile", help="Describe a file without importing it")
    prof.add_argument("--input", required=True)
    prof.add_argument("--output")

    mapper = sub.add_parser("map", help="Propose a mapping for review")
    mapper.add_argument("--input", required=True)
    mapper.add_argument("--target", required=True, choices=["suppliers", "spend"])
    mapper.add_argument("--source-system", default="SRC")
    mapper.add_argument("--date-order", choices=["dmy", "mdy", "ymd"])
    mapper.add_argument("--decimal-style", choices=["point", "comma", "plain"])
    mapper.add_argument("--column", action="append", default=[], metavar="FIELD=SOURCE", help="Explicitly confirm or override a column mapping")
    mapper.add_argument("--output", required=True)

    prep = sub.add_parser("prepare", help="Emit canonical suppliers.csv and spend.csv")
    prep.add_argument("--suppliers", required=True)
    prep.add_argument("--spend", required=True)
    prep.add_argument("--supplier-mapping")
    prep.add_argument("--spend-mapping")
    prep.add_argument("--source-system", default="SRC")
    prep.add_argument("--date-order", choices=["dmy", "mdy", "ymd"])
    prep.add_argument("--decimal-style", choices=["point", "comma", "plain"])
    prep.add_argument("--invoice-key-scope", choices=["dataset", "supplier"],
                      help="Use 'supplier' when invoice numbers repeat across vendors")
    prep.add_argument("--output-dir", required=True)

    args = parser.parse_args(argv)
    try:
        if args.command == "profile":
            report = profile_source(_load(args.input), Path(args.input).name).report()
            text = json.dumps(report, indent=2)
            if args.output:
                Path(args.output).parent.mkdir(parents=True, exist_ok=True)
                Path(args.output).write_text(text, encoding="utf-8")
            print(text)
            return 0

        if args.command == "map":
            data = _load(args.input)
            profile = profile_source(data, Path(args.input).name)
            extra = {"date_order": args.date_order} if args.date_order else {}
            if args.decimal_style:
                extra["decimal_style"] = args.decimal_style
            if args.column:
                try:
                    extra["column_overrides"] = dict(value.split("=", 1) for value in args.column)
                except ValueError as exc:
                    raise IngestError("--column must be FIELD=SOURCE") from exc
            mapping = infer_mapping(profile, args.target, args.source_system, data, **extra)
            Path(args.output).parent.mkdir(parents=True, exist_ok=True)
            Path(args.output).write_text(mapping.to_json(), encoding="utf-8")
            print(json.dumps({"mapping": args.output, "mapping_id": mapping.mapping_id,
                              "columns": mapping.columns, "notes": profile.notes}, indent=2))
            return 0

        supplier_mapping = SourceMapping.from_json(Path(args.supplier_mapping).read_text(encoding="utf-8")) if args.supplier_mapping else None
        spend_mapping = SourceMapping.from_json(Path(args.spend_mapping).read_text(encoding="utf-8")) if args.spend_mapping else None
        spend_options = {}
        if args.date_order:
            spend_options["date_order"] = args.date_order
        if args.decimal_style:
            spend_options["decimal_style"] = args.decimal_style
        if args.invoice_key_scope:
            spend_options["invoice_key_scope"] = args.invoice_key_scope
        inference = {"spend": spend_options} if spend_options else {}
        prepared = prepare_dataset(args.suppliers, args.spend, supplier_mapping, spend_mapping,
                                   args.source_system, Path(args.suppliers).name,
                                   Path(args.spend).name, **inference)
        out = Path(args.output_dir)
        out.mkdir(parents=True, exist_ok=True)
        (out / "suppliers.csv").write_text(prepared.suppliers_csv, encoding="utf-8")
        (out / "spend.csv").write_text(prepared.spend_csv, encoding="utf-8")
        (out / "extras.json").write_text(json.dumps(prepared.extras, indent=2), encoding="utf-8")
        for target, mapping in prepared.manifest["mappings"].items():
            (out / f"{target}-mapping.json").write_text(json.dumps(mapping, indent=2), encoding="utf-8")
        # Completion evidence is written last; a partially updated CSV cannot
        # validate against a previous manifest.
        (out / "manifest.json").write_text(json.dumps(prepared.manifest, indent=2), encoding="utf-8")
        print(json.dumps({"output_dir": str(out), "manifest_id": prepared.manifest["manifest_id"],
                          "suppliers": prepared.manifest["suppliers"],
                          "spend": prepared.manifest["spend"]}, indent=2))
        return 0
    except (IngestError, OSError, ValueError) as exc:
        parser.exit(2, f"Ingest error: {exc}\n")


if __name__ == "__main__":
    raise SystemExit(main())
